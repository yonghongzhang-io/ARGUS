"""Build the annotation kit for the human pilot: one workbook per annotator.

    python3 annotation/human_pilot/build_kit.py

Writes data/annotations/human_pilot/to_send/ (private, not tracked): ARGUS_annotation_A.xlsx,
ARGUS_annotation_B.xlsx, GUIDELINE.md (+ .docx) and the five PDFs. Sheets are protected so that only
the answer cells can be edited. The workbooks contain the rubric text
only: no system output, no earlier labels, nothing about what any result would imply.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = ROOT / "data" / "annotations" / "human_pilot" / "to_send"
PDFS = ROOT.parent / "CAUSALVERIFY" / "v11" / "pdfs-corpus" / "raw"
PAPERS = ["paper_01", "paper_03", "paper_07", "paper_08", "paper_10"]
COLS = [("#", 4), ("dimension", 20), ("assumption", 34), ("evidence a credible paper would report", 38),
        ("applicability", 15), ("reported", 11), ("evidence location", 26), ("verbatim quote (optional)", 40),
        ("risk", 10), ("confidence", 11), ("rationale", 52)]
FILL_IN = PatternFill("solid", fgColor="FFF9C4")
HEAD = PatternFill("solid", fgColor="1F3A5F")
GREY = PatternFill("solid", fgColor="F2F4F7")
THIN = Side(style="thin", color="C9CED6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)

RULES = [
    "Read each PDF yourself. Do not use ChatGPT, Claude, Copilot or any other AI tool to read the paper, find passages, "
    "choose labels or write rationales. Searching inside the PDF (Ctrl/Cmd-F) is fine.",
    "Work alone. Do not discuss these papers or your labels with the other annotator until both workbooks are returned.",
    "Do not look at any ARGUS output for these papers. There is no expected answer.",
    "Judge what the paper reports. A check that may have been run but is not reported counts as not reported.",
    "Fill only the yellow cells. One sheet per paper, eleven rows each; then the Sign-off sheet. About one hour per paper.",
]
LEGEND = [
    ("applicability", "conventional = applies as usually defined for DID; analogue = the design is not a canonical treated-vs-control "
                      "DID but an equivalent assumption must hold (judge that analogue and name it in the rationale); "
                      "not applicable = no counterpart at all (leave risk empty, explain in the rationale)."),
    ("reported", "yes / no / unclear: does the paper report any evidence or check bearing on the assumption or its analogue?"),
    ("evidence location", "Section, table, figure or page you relied on; 'none found' if none."),
    ("verbatim quote", "Optional. One to three sentences copied exactly from the PDF."),
    ("risk", "low = a relevant check is reported and its result directly supports the assumption. medium = partial support with a "
             "named, unresolved risk. high = barely addressed with no diagnostic evidence, or the evidence itself points to a violation."),
    ("confidence", "1 to 5: how sure you are of the risk label (5 = very sure)."),
    ("rationale", "One or two sentences in your own words. Say so if experts could reasonably disagree."),
]
EXAMPLE = ["ex.", "Parallel trends", "(example row, from a hypothetical paper that is not one of the five)", "",
           "conventional", "yes", "Section 5.1; Figure 3", "\"Pre-period coefficients are small and jointly insignificant (p = 0.62).\"",
           "low", 4, "Event study with six leads is shown; leads are flat and jointly insignificant, which directly supports the assumption."]


def titles() -> dict[str, str]:
    import csv
    with (ROOT / "experiments" / "real_papers" / "corpus_manifest.csv").open(encoding="utf-8") as fh:
        return {r["paper_id"]: r["title_as_indexed"] for r in csv.DictReader(fh)}


def lock(ws) -> None:
    """Only the yellow cells can be edited, so the layout the receiving script reads cannot be broken.
    No password: the point is to prevent accidents, and rows and columns can still be resized."""
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD, Font(name="Arial", size=10, bold=True, color="FFFFFF"), WRAP, BOX


def instructions(wb: Workbook, who: str) -> None:
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120
    ws["A1"] = f"ARGUS annotation pilot: workbook for annotator {who}"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    r = 3
    ws.cell(row=r, column=1, value="Ground rules").font = BOLD
    for i, rule in enumerate(RULES, 1):
        r += 1
        ws.cell(row=r, column=1, value=f"{i}.").font = FONT
        c = ws.cell(row=r, column=2, value=rule); c.font, c.alignment = FONT, WRAP
    r += 2
    ws.cell(row=r, column=1, value="Columns to fill").font = BOLD
    for name, text in LEGEND:
        r += 1
        ws.cell(row=r, column=1, value=name).font = BOLD
        c = ws.cell(row=r, column=2, value=text); c.font, c.alignment = FONT, WRAP
    r += 2
    ws.cell(row=r, column=1, value="More detail").font = BOLD
    c = ws.cell(row=r, column=2, value="Borderline conventions and the full definitions are in GUIDELINE.md, sent with this workbook."); c.font = FONT


def paper_sheet(wb: Workbook, pid: str, title: str, dims: list[dict]) -> None:
    ws = wb.create_sheet(pid)
    ws["A1"] = f"{pid}.pdf"; ws["A1"].font = Font(name="Arial", size=12, bold=True)
    ws["C1"] = title; ws["C1"].font = FONT
    ws["A2"] = "Fill the yellow cells. Row 5 is an example of the expected format and is not part of the task."
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="555555")
    for c, (name, width) in enumerate(COLS, 1):
        ws.cell(row=4, column=c, value=name)
        ws.column_dimensions[ws.cell(row=4, column=c).column_letter].width = width
    style_header(ws, 4, len(COLS))
    for c, v in enumerate(EXAMPLE, 1):
        cell = ws.cell(row=5, column=c, value=v)
        cell.font, cell.alignment, cell.border, cell.fill = Font(name="Arial", size=9, italic=True, color="555555"), WRAP, BOX, GREY
    for i, d in enumerate(dims, 1):
        row = 5 + i
        evidence = "; ".join(d["evidence"]) if isinstance(d["evidence"], list) else d["evidence"]
        for c, v in enumerate([i, d["name"], d["assumption"], evidence], 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font, cell.alignment, cell.border, cell.fill = (BOLD if c == 2 else FONT), WRAP, BOX, GREY
        for c in range(5, len(COLS) + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill, cell.font, cell.alignment, cell.border = FILL_IN, FONT, WRAP, BOX
            cell.protection = Protection(locked=False)
        ws.row_dimensions[row].height = 78
    ws.row_dimensions[5].height = 52
    for col, formula, msg in (("E", '"conventional,analogue,not applicable"', "conventional / analogue / not applicable"),
                              ("F", '"yes,no,unclear"', "yes / no / unclear"), ("I", '"low,medium,high"', "low / medium / high"),
                              ("J", '"1,2,3,4,5"', "1 to 5")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorTitle="Not an allowed value", error=f"Please choose: {msg}")
        ws.add_data_validation(dv)
        dv.add(f"{col}6:{col}16")
    ws.freeze_panes = "C5"
    lock(ws)


def signoff(wb: Workbook, who: str) -> None:
    ws = wb.create_sheet("Sign-off")
    ws.column_dimensions["A"].width = 100
    ws.column_dimensions["B"].width = 34
    ws["A1"] = f"Sign-off, annotator {who}"; ws["A1"].font = Font(name="Arial", size=14, bold=True)
    rows = ["Your name", "Date(s) on which you annotated", "Approximate total hours",
            "I read the five PDFs myself (yes / no)",
            "I used no AI tool to read the papers, find evidence, choose labels or write rationales (yes / no)",
            "I did not see any ARGUS output for these papers (yes / no)",
            "I did not discuss these labels with the other annotator before returning this workbook (yes / no)",
            "Anything we should know (optional)"]
    for i, text in enumerate(rows, 3):
        ws.cell(row=i, column=1, value=text).font = FONT
        c = ws.cell(row=i, column=2); c.fill, c.border, c.font = FILL_IN, BOX, FONT
        c.protection = Protection(locked=False)
    dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B6:B9")
    lock(ws)


def main() -> None:
    dims = yaml.safe_load((ROOT / "config" / "identification_dimensions.yaml").read_text(encoding="utf-8"))
    dims = dims["dimensions"] if isinstance(dims, dict) else dims
    assert len(dims) == 11
    OUT.mkdir(parents=True, exist_ok=True)
    t = titles()
    for who in "AB":
        wb = Workbook()
        instructions(wb, who)
        for pid in PAPERS:
            paper_sheet(wb, pid, t[pid], dims)
        signoff(wb, who)
        wb.save(OUT / f"ARGUS_annotation_{who}.xlsx")
    shutil.copy(HERE / "GUIDELINE.md", OUT / "GUIDELINE.md")
    if shutil.which("pandoc"):  # a Word copy is easier to open than Markdown
        import subprocess
        subprocess.run(["pandoc", str(HERE / "GUIDELINE.md"), "-o", str(OUT / "GUIDELINE.docx")], check=True)
    for pid in PAPERS:
        shutil.copy(PDFS / f"{pid}.pdf", OUT / f"{pid}.pdf")
    print("wrote", OUT)


if __name__ == "__main__":
    main()
